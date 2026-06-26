import io
from typing import List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def generate_excel(headers: List[str], rows: List[List[Any]], sheet_name: str = "Sheet1") -> bytes:
    """
    Generates an Excel file as bytes given a list of headers and rows.
    Automatically applies basic styling to the header row and auto-fits column widths.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Write headers
    ws.append(headers)

    # Style headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Write rows
    for row_data in rows:
        ws.append(row_data)

    # Auto-fit column widths
    for col_idx, header in enumerate(headers, 1):
        max_length = len(str(header))
        for row_idx in range(2, len(rows) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value is not None:
                max_length = max(max_length, len(str(cell_value)))
        
        # Add a little extra padding
        adjusted_width = (max_length + 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    # Save to memory stream
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
